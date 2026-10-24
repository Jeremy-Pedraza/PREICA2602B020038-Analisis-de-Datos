# -*- coding: utf-8 -*-
"""Genera el documento Word de la entrega AE2 con python-docx.

Abre la plantilla institucional ``AE2/documento/plantilla_iudigital.docx``
(conserva encabezado con logo, márgenes y estilos), limpia el cuerpo y
construye el documento completo de la AE2.

REGLA DURA: ninguna cifra se escribe a mano. Todas las dimensiones, tipos,
nulos, frecuencias y estadísticas se leen en tiempo de ejecución desde
``AE2/resultados/churn_crudo.csv``, ``AE2/resultados/churn_transformado.csv``
y ``AE2/resultados/resumen_dataframe.txt``.

Reproducible: volver a ejecutarlo regenera el .docx.

Uso:
    python AE2/src/generar_documento.py
"""

from __future__ import annotations

import pathlib

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

# --------------------------------------------------------------------------
# Rutas
# --------------------------------------------------------------------------
AE2 = pathlib.Path(__file__).resolve().parent.parent
RESULTADOS = AE2 / "resultados"
DOCUMENTO = AE2 / "documento"
RAIZ = AE2.parent

PLANTILLA = DOCUMENTO / "plantilla_iudigital.docx"
SALIDA = DOCUMENTO / "Pedraza_Jeremy_AE2_Analitica_de_Datos.docx"

CSV_CRUDO = RESULTADOS / "churn_crudo.csv"
CSV_TRANS = RESULTADOS / "churn_transformado.csv"
TXT_RESUMEN = RESULTADOS / "resumen_dataframe.txt"
REQUIREMENTS = RAIZ / "requirements.txt"

# --------------------------------------------------------------------------
# Datos institucionales (literales)
# --------------------------------------------------------------------------
TITULO = (
    "Caracterización de la cancelación del servicio en clientes de "
    "telecomunicaciones: análisis descriptivo del dataset Telco Customer Churn"
)
INTEGRANTE = "Jeremy Iván Pedraza Hernández"
GRUPO = "Grupo 1"
DOCENTE = "Andrés Felipe Callejas Jaramillo"
ASIGNATURA = "Proyecto Integrado III - Analítica de Datos"
CODIGO = "PREICA2602B020038"
PROGRAMA = "Ingeniería de Software y Datos"
FACULTAD = "Facultad de Ingeniería y Ciencias Agropecuarias"
INSTITUCION = "Institución Universitaria Digital de Antioquia."
ANIO = "2026"
REPO = "https://github.com/Jeremy-Pedraza/PREICA2602B020038-Analisis-de-Datos"
DATASET_ID = "scikit-learn/churn-prediction"
DATASET_URL = "https://huggingface.co/datasets/scikit-learn/churn-prediction"

# Diccionario de significados conocidos del dataset Telco Customer Churn.
# Lo que no esté documentado se marca como "Pendiente de verificar".
SIGNIFICADOS = {
    "customer_id": "Identificador único del cliente.",
    "gender": "Género del cliente.",
    "senior_citizen": "Indica si el cliente es adulto mayor (1) o no (0).",
    "partner": "Indica si el cliente tiene pareja.",
    "dependents": "Indica si el cliente tiene personas a cargo.",
    "tenure": "Antigüedad del cliente en meses con la compañía.",
    "phone_service": "Indica si el cliente tiene servicio telefónico.",
    "multiple_lines": "Indica si el cliente tiene múltiples líneas telefónicas.",
    "internet_service": "Tipo de servicio de internet contratado.",
    "online_security": "Indica si el cliente tiene seguridad en línea contratada.",
    "online_backup": "Indica si el cliente tiene copia de seguridad en línea.",
    "device_protection": "Indica si el cliente tiene protección de dispositivos.",
    "tech_support": "Indica si el cliente tiene soporte técnico contratado.",
    "streaming_tv": "Indica si el cliente tiene servicio de televisión por streaming.",
    "streaming_movies": "Indica si el cliente tiene servicio de películas por streaming.",
    "contract": "Tipo de contrato del cliente (mensual, anual o bianual).",
    "paperless_billing": "Indica si el cliente usa facturación electrónica.",
    "payment_method": "Método de pago utilizado por el cliente.",
    "monthly_charges": "Cargo mensual facturado al cliente (unidad monetaria).",
    "total_charges": "Cargo total acumulado facturado al cliente (unidad monetaria).",
    "churn": "Variable objetivo: indica si el cliente canceló el servicio.",
    "cargo_promedio_mensual": (
        "Variable derivada: total_charges / tenure, cargo promedio por mes de "
        "antigüedad (unidad monetaria)."
    ),
    "segmento_antiguedad": (
        "Variable derivada: segmentación de tenure en rangos "
        "(0-12, 13-24, 25-48, 49+ meses)."
    ),
}


# --------------------------------------------------------------------------
# Utilidades de documento
# --------------------------------------------------------------------------
def limpiar_cuerpo(doc: Document) -> None:
    """Elimina todo el contenido del cuerpo conservando el sectPr final."""
    body = doc.element.body
    sectPr = body.find(qn("w:sectPr"))
    for child in list(body):
        if child is not sectPr:
            body.remove(child)


def p_centrado(doc: Document, texto: str, negrita: bool = False) -> None:
    """Agrega un párrafo centrado (estilo Normal), opcionalmente en negrita."""
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(texto)
    run.bold = negrita


def add_field(paragraph, instr: str) -> None:
    """Inserta un campo de Word (p. ej. TOC) como XML dentro del párrafo."""
    r = paragraph.add_run()
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t")
    txt.text = "Actualizar este campo en Word con F9"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (fld, it, sep, txt, end):
        r._r.append(el)


def h1(doc: Document, texto: str):
    return doc.add_heading(texto, level=1)


def h2(doc: Document, texto: str):
    return doc.add_heading(texto, level=2)


def parrafo(doc: Document, texto: str):
    return doc.add_paragraph(texto, style="Normal")


def item(doc: Document, texto: str):
    p = doc.add_paragraph(style="List Paragraph")
    p.add_run("- " + texto)
    return p


def agregar_tabla(doc: Document, encabezados, filas, numero, titulo, fuente):
    """Crea una tabla con estilo institucional, encabezado en negrita y pie Caption."""
    tabla = doc.add_table(rows=1, cols=len(encabezados))
    try:
        tabla.style = "Grid Table 2 Accent 4"
    except KeyError:
        tabla.style = "Table Grid"
    hdr = tabla.rows[0].cells
    for i, texto in enumerate(encabezados):
        hdr[i].text = ""
        run = hdr[i].paragraphs[0].add_run(str(texto))
        run.bold = True
    for fila in filas:
        celdas = tabla.add_row().cells
        for i, valor in enumerate(fila):
            celdas[i].text = str(valor)
    cap = doc.add_paragraph(style="Caption")
    cap.add_run(f"Tabla {numero}. {titulo} ").bold = True
    cap.add_run(f"Fuente: {fuente}")
    return tabla


def parrafo_mono(doc: Document, lineas) -> None:
    """Agrega cada línea como párrafo con fuente monoespaciada."""
    for linea in lineas:
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(linea if linea else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(7.5)


# --------------------------------------------------------------------------
# Lectura de datos reales
# --------------------------------------------------------------------------
def leer_fecha_consulta() -> str:
    for linea in TXT_RESUMEN.read_text(encoding="utf-8").splitlines():
        if "fecha_consulta" in linea:
            return linea.split(":", 1)[1].strip()
    return "Pendiente de verificar"


def leer_versiones_requirements() -> list[tuple[str, str]]:
    pares = []
    for linea in REQUIREMENTS.read_text(encoding="utf-8").splitlines():
        linea = linea.strip()
        if not linea or linea.startswith("#"):
            continue
        for sep in (">=", "==", "<=", "~="):
            if sep in linea:
                nombre, version = linea.split(sep, 1)
                pares.append((nombre.strip(), f"{sep}{version.strip()}"))
                break
        else:
            pares.append((linea, "sin fijar"))
    return pares


# --------------------------------------------------------------------------
# Construcción del documento
# --------------------------------------------------------------------------
def construir() -> None:
    df_crudo = pd.read_csv(CSV_CRUDO)
    df = pd.read_csv(CSV_TRANS)
    resumen_txt = TXT_RESUMEN.read_text(encoding="utf-8").splitlines()

    # --- Métricas derivadas de los datos (no se inventan) ---
    n_crudo_f, n_crudo_c = df_crudo.shape
    n_f, n_c = df.shape
    fecha = leer_fecha_consulta()

    # Nulos de total_charges en el crudo (texto con vacíos -> coerce).
    tc_crudo = pd.to_numeric(df_crudo["TotalCharges"], errors="coerce")
    nulos_tc = int(tc_crudo.isna().sum())
    pct_nulos_tc = nulos_tc / n_crudo_f * 100

    churn_vc = df["churn"].value_counts()
    churn_si = int(churn_vc.get("Yes", 0))
    churn_no = int(churn_vc.get("No", 0))
    pct_churn = churn_si / n_f * 100

    contract_vc = df["contract"].value_counts()
    internet_vc = df["internet_service"].value_counts()
    pago_vc = df["payment_method"].value_counts()
    desc = df[["tenure", "monthly_charges", "total_charges"]].describe()

    doc = Document(str(PLANTILLA))
    limpiar_cuerpo(doc)

    # =========================== PORTADA ===========================
    p_centrado(doc, TITULO, negrita=True)
    p_centrado(doc, "")
    p_centrado(doc, "")
    p_centrado(doc, "Integrantes", negrita=True)
    p_centrado(doc, GRUPO, negrita=True)
    p_centrado(doc, INTEGRANTE, negrita=False)
    p_centrado(doc, "")
    p_centrado(doc, "Docente", negrita=True)
    p_centrado(doc, DOCENTE, negrita=True)
    p_centrado(doc, "")
    p_centrado(doc, ASIGNATURA, negrita=True)
    p_centrado(doc, CODIGO, negrita=False)
    p_centrado(doc, "")
    p_centrado(doc, "")
    p_centrado(doc, "")
    p_centrado(doc, PROGRAMA, negrita=True)
    p_centrado(doc, FACULTAD, negrita=True)
    p_centrado(doc, INSTITUCION, negrita=True)
    p_centrado(doc, ANIO, negrita=True)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # =========================== CONTENIDO (TOC) ===========================
    doc.add_paragraph("Contenido", style="TOC Heading")
    add_field(doc.add_paragraph(), 'TOC \\o "1-3" \\h \\z \\u')
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    doc.add_paragraph("Lista de tablas", style="TOC Heading")
    add_field(doc.add_paragraph(), 'TOC \\h \\z \\c "Tabla"')
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # =========================== 1. INTRODUCCIÓN ===========================
    h1(doc, "1. Introducción")
    parrafo(
        doc,
        "El sector de las telecomunicaciones opera en un mercado altamente "
        "competitivo en el que la retención de clientes es un factor determinante "
        "de la sostenibilidad del negocio. La cancelación del servicio, conocida "
        "como rotación o churn, representa la pérdida de clientes activos y, con "
        "ella, de los ingresos recurrentes asociados. En este contexto, el análisis "
        "descriptivo de los datos de clientes permite caracterizar los perfiles que "
        "concentran la cancelación y orientar decisiones comerciales informadas.",
    )
    parrafo(
        doc,
        f"El presente documento corresponde a la primera etapa del proyecto "
        f"integrador y utiliza el conjunto de datos público «{DATASET_ID}» (Telco "
        f"Customer Churn), disponible en Hugging Face. El dataset reúne "
        f"{n_crudo_f} registros de clientes y {n_crudo_c} variables originales que "
        f"describen su perfil, los servicios contratados y su facturación, además de "
        f"la variable objetivo churn.",
    )
    parrafo(
        doc,
        "El propósito de esta etapa es reconocer el dataset, evaluar su calidad, "
        "preparar las variables y documentar la necesidad analítica que guiará las "
        "etapas siguientes. El documento se organiza en el planteamiento de la "
        "necesidad, la justificación, los objetivos, el marco teórico inicial, los "
        "materiales y métodos, los resultados del reconocimiento del dataset, las "
        "conclusiones y las recomendaciones.",
    )
    parrafo(
        doc,
        "Este trabajo constituye la primera de tres etapas del proyecto integrador; "
        "las etapas posteriores abordarán el análisis exploratorio en profundidad y "
        "el modelado orientado a la predicción de la cancelación.",
    )

    # =================== 2. PLANTEAMIENTO DE LA NECESIDAD ===================
    h1(doc, "2. Planteamiento de la necesidad")
    h2(doc, "2.1 Contexto y beneficiario")
    parrafo(
        doc,
        "El beneficiario de este análisis es el área comercial y de retención de "
        "clientes de una compañía de telecomunicaciones. Esta área enfrenta "
        "permanentemente la decisión de dónde concentrar sus esfuerzos de retención, "
        "ya que no todos los clientes presentan el mismo riesgo de cancelación. "
        "Disponer de una caracterización de los clientes que cancelan frente a los "
        "que permanecen apoya la asignación eficiente de recursos de retención.",
    )
    h2(doc, "2.2 Necesidad de información")
    parrafo(
        doc,
        "Se requiere caracterizar qué perfiles de clientes concentran la cancelación "
        "del servicio, considerando variables contractuales (contract, "
        "paperless_billing, payment_method), de servicios (internet_service, "
        "online_security, tech_support, streaming_tv, entre otras) y de facturación "
        "(monthly_charges, total_charges, tenure). Se declara de forma explícita que "
        "esta es una necesidad propuesta con fines académicos y no un problema "
        "observado y comprobado en una organización real.",
    )
    h2(doc, "2.3 Pregunta de análisis")
    parrafo(
        doc,
        "¿Qué características contractuales, de servicio y de facturación distinguen "
        "a los clientes que cancelan el servicio frente a los que permanecen? Esta "
        "pregunta puede responderse con las variables disponibles en el dataset: la "
        "variable objetivo churn, las variables contractuales contract, "
        "paperless_billing y payment_method, las variables de servicios "
        "internet_service, online_security, online_backup, device_protection, "
        "tech_support, streaming_tv y streaming_movies, y las variables de "
        "facturación tenure, monthly_charges y total_charges.",
    )
    h2(doc, "2.4 Escenario")
    parrafo(
        doc,
        "El escenario es hipotético. El dataset es un conjunto de referencia público "
        "y no corresponde a una empresa colombiana ni a una organización específica; "
        "se emplea como caso de estudio para ejercitar el proceso de analítica "
        "descriptiva dentro de la asignatura.",
    )

    # =========================== 3. JUSTIFICACIÓN ===========================
    h1(doc, "3. Justificación")
    parrafo(
        doc,
        "El costo de adquirir un cliente nuevo suele ser considerablemente mayor que "
        "el de retener a un cliente existente, por lo que anticipar y comprender la "
        "cancelación tiene un impacto económico directo. Un análisis descriptivo "
        "sólido es el punto de partida para priorizar acciones de retención y para "
        "fundamentar modelos predictivos posteriores.",
    )
    parrafo(
        doc,
        "El dataset seleccionado es pertinente porque combina variables "
        "contractuales, de servicios contratados y de facturación junto con la "
        "variable objetivo churn, lo que permite abordar la necesidad planteada sin "
        "requerir fuentes adicionales. Académicamente, el ejercicio consolida las "
        "competencias de carga, exploración, preparación y documentación de datos "
        "propias de la asignatura Proyecto Integrado III - Analítica de Datos.",
    )

    # =========================== 4. OBJETIVOS ===========================
    h1(doc, "4. Objetivos")
    h2(doc, "4.1 Objetivo general")
    parrafo(
        doc,
        "Caracterizar, mediante análisis descriptivo, las diferencias contractuales, "
        "de servicio y de facturación entre los clientes que cancelan el servicio y "
        "los que permanecen, a partir del dataset Telco Customer Churn.",
    )
    h2(doc, "4.2 Objetivos específicos")
    for texto in [
        "Obtener y cargar el dataset desde Hugging Face en un DataFrame de pandas.",
        "Reconocer la estructura y la calidad del dataset (dimensiones, tipos de "
        "datos, valores nulos y duplicados).",
        "Preparar y documentar las variables mediante transformaciones de limpieza, "
        "tipado, derivación y segmentación.",
        "Caracterizar y comparar de forma descriptiva los grupos de permanencia y "
        "cancelación según las variables disponibles.",
    ]:
        item(doc, texto)

    # =========================== 5. MARCO TEÓRICO ===========================
    h1(doc, "5. Marco teórico inicial")
    h2(doc, "5.1 Analítica descriptiva")
    parrafo(
        doc,
        "La analítica descriptiva busca resumir y describir las características "
        "principales de un conjunto de datos para responder qué ha ocurrido, "
        "mediante medidas de resumen, frecuencias y distribuciones (Provost & "
        "Fawcett, 2013). Se diferencia de la analítica predictiva, que estima "
        "resultados futuros, y de la prescriptiva, que recomienda acciones; la "
        "descriptiva constituye el fundamento sobre el que se construyen ambas.",
    )
    h2(doc, "5.2 Rotación de clientes (churn)")
    parrafo(
        doc,
        "La rotación de clientes o churn es la pérdida de clientes durante un "
        "periodo determinado. La tasa de cancelación es un indicador clave de la "
        "gestión comercial, pues afecta directamente los ingresos recurrentes y el "
        "valor de vida del cliente. Su análisis permite identificar factores "
        "asociados al abandono y diseñar estrategias de retención (Verbeke et al., "
        "2012).",
    )
    h2(doc, "5.3 Calidad de datos")
    parrafo(
        doc,
        "La calidad de los datos condiciona la validez de cualquier conclusión. "
        "Aspectos como los valores nulos, el tipado correcto de las variables y la "
        "presencia de registros duplicados deben evaluarse y tratarse antes del "
        "análisis, ya que de lo contrario pueden sesgar los resultados (Han, "
        "Kamber & Pei, 2011).",
    )
    h2(doc, "5.4 Herramientas")
    parrafo(
        doc,
        "Hugging Face Datasets es un repositorio y una librería que facilita el "
        "acceso reproducible a conjuntos de datos públicos (Lhoest et al., 2021). "
        "pandas es la librería estándar de Python para la manipulación de datos "
        "tabulares, con estructuras como el DataFrame que permiten cargar, "
        "transformar y resumir información de manera eficiente (McKinney, 2010).",
    )

    # =================== 6. MATERIALES Y MÉTODOS ===================
    h1(doc, "6. Materiales y métodos (metodología de la etapa 1)")

    h2(doc, "6.1 Ficha del dataset")
    ficha = [
        ["Nombre", "Telco Customer Churn"],
        ["Identificador del repositorio", DATASET_ID],
        ["Enlace directo", DATASET_URL],
        ["Propietario", "scikit-learn (organización en Hugging Face)"],
        ["Licencia y condiciones de uso", "Según ficha del dataset en Hugging Face"],
        ["Modalidad", "Tabular (CSV)"],
        ["Tamaño", f"{n_crudo_f} filas x {n_crudo_c} columnas (crudo)"],
        ["Split utilizado", "train"],
        ["Revisión", "No fijada (última disponible)"],
        ["Fecha de consulta", fecha],
        ["Naturaleza de los datos", "De referencia pública; según ficha"],
    ]
    agregar_tabla(
        doc, ["Campo", "Valor"], ficha, 1, "Ficha del dataset.",
        "elaboración propia a partir de la ficha del dataset en Hugging Face.",
    )

    h2(doc, "6.2 Justificación de la selección")
    parrafo(
        doc,
        "Frente a alternativas como datasets de clasificación genérica (por ejemplo "
        "Iris, usado en la EA1) o datasets sin variable objetivo de negocio, Telco "
        "Customer Churn se seleccionó porque plantea un problema real de negocio, "
        "incluye una variable objetivo (churn) y combina atributos contractuales, de "
        "servicios y de facturación, lo que lo hace idóneo para la necesidad "
        "planteada.",
    )

    h2(doc, "6.3 Unidad de análisis")
    parrafo(
        doc,
        "La unidad de análisis es el cliente: cada fila del DataFrame representa un "
        "cliente del servicio de telecomunicaciones, identificado por customer_id, "
        "con sus atributos de perfil, servicios contratados, facturación y su estado "
        "de cancelación (churn).",
    )

    h2(doc, "6.4 Diccionario inicial de variables")
    filas_dicc = []
    for col in df.columns:
        dtype = str(df[col].dtype)
        serie = df[col]
        if dtype == "object" or dtype == "category":
            cats = serie.dropna().unique().tolist()
            if len(cats) <= 6:
                detalle = "Categorías: " + ", ".join(str(c) for c in cats)
            else:
                detalle = f"{len(cats)} valores distintos (p. ej. identificador)"
        else:
            detalle = f"Numérica (mín={serie.min()}, máx={serie.max()})"
        significado = SIGNIFICADOS.get(col, "Pendiente de verificar")
        filas_dicc.append([col, significado, dtype, detalle])
    agregar_tabla(
        doc,
        ["Variable", "Significado", "Tipo", "Categorías / unidad"],
        filas_dicc, 2, "Diccionario inicial de variables.",
        "elaboración propia a partir del DataFrame transformado.",
    )

    h2(doc, "6.5 Herramientas y entorno")
    parrafo(
        doc,
        "El desarrollo se realizó en Python con Jupyter Notebook, control de "
        "versiones con Git y alojamiento en GitHub. Las librerías principales y sus "
        "versiones mínimas, según requirements.txt, se presentan en la Tabla 3.",
    )
    filas_req = [[n, v] for n, v in leer_versiones_requirements()]
    agregar_tabla(
        doc, ["Librería / herramienta", "Versión"], filas_req, 3,
        "Herramientas y versiones del entorno.",
        "elaboración propia a partir de requirements.txt.",
    )

    h2(doc, "6.6 Procedimiento ejecutado")
    parrafo(
        doc,
        "El procedimiento de la etapa 1 comprendió: (1) la carga del dataset "
        "mediante la clase HFDatasetLoader, que se conecta a Hugging Face y devuelve "
        "un DataFrame; (2) la exploración descriptiva inicial; (3) cinco "
        "transformaciones de preparación; y (4) la generación automática del resumen "
        "con la clase DatasetSummarizer.",
    )
    for texto in [
        "T1 - Renombrado de columnas a snake_case para estandarizar su notación.",
        "T2 - Conversión de total_charges de texto a numérico con "
        "pd.to_numeric(errors='coerce'), dado que contenía cadenas vacías.",
        "T3 - Tratamiento de los nulos resultantes de total_charges, imputados con "
        "0.0 por corresponder a clientes con tenure = 0.",
        "T4 - Creación de la variable derivada cargo_promedio_mensual = "
        "total_charges / tenure, protegiendo la división por cero.",
        "T5 - Segmentación de la antigüedad (tenure) en rangos mediante pd.cut, "
        "generando segmento_antiguedad.",
    ]:
        item(doc, texto)
    parrafo(
        doc,
        "Archivos generados (rutas relativas a la raíz del repositorio): "
        "AE2/resultados/churn_crudo.csv, AE2/resultados/churn_transformado.csv y "
        "AE2/resultados/resumen_dataframe.txt. El notebook de referencia es "
        "AE2/notebooks/ae2_exploracion_descriptiva.ipynb.",
    )

    h2(doc, "6.7 Alcance y limitaciones")
    parrafo(
        doc,
        "El alcance de esta etapa es descriptivo. La unidad de análisis es el "
        "cliente y la cobertura corresponde a los registros del split train. Los "
        "datos no incluyen marca temporal, por lo que no permiten establecer "
        "tendencias en el tiempo. De forma explícita, los datos no permiten inferir "
        "causalidad, ni extrapolar los resultados a otro mercado o país, ni afirmar "
        "hallazgos predictivos en esta etapa.",
    )

    h2(doc, "6.8 Diseño inicial del análisis")
    parrafo(
        doc,
        "Las variables relevantes para la pregunta son churn (objetivo), contract, "
        "payment_method, internet_service, tenure, monthly_charges y total_charges. "
        "Los pasos previstos de la etapa descriptiva son: describir la distribución "
        "de la variable objetivo, comparar las frecuencias de las variables "
        "contractuales y de servicio entre clientes que cancelan y permanecen, y "
        "resumir las variables de facturación por grupo. Cada paso se relaciona "
        "directamente con la pregunta de análisis al contrastar los dos grupos de "
        "interés.",
    )

    # =========================== 7. RESULTADOS ===========================
    h1(doc, "7. Resultados - reconocimiento del dataset")

    h2(doc, "7.1 Dimensiones y estructura")
    parrafo(
        doc,
        f"El dataset crudo contiene {n_crudo_f} filas y {n_crudo_c} columnas. Tras "
        f"la preparación, el DataFrame transformado conserva {n_f} filas y amplía a "
        f"{n_c} columnas por la incorporación de las variables derivadas "
        f"cargo_promedio_mensual y segmento_antiguedad. No se eliminaron registros: "
        f"no se detectaron filas duplicadas.",
    )

    h2(doc, "7.2 Tipos de datos y valores nulos")
    parrafo(
        doc,
        f"La variable total_charges venía originalmente como texto con cadenas "
        f"vacías; al convertirla a numérica se identificaron {nulos_tc} valores "
        f"nulos ({pct_nulos_tc:.2f}% del total), todos correspondientes a clientes "
        f"con tenure = 0, que fueron imputados con 0.0. Tras la preparación, el "
        f"DataFrame transformado no presenta valores nulos. La Tabla 4 resume el "
        f"tipo de dato y los nulos por columna en el DataFrame transformado.",
    )
    filas_nulos = []
    total = len(df)
    for col in df.columns:
        n_nulos = int(df[col].isna().sum())
        pct = n_nulos / total * 100
        filas_nulos.append([col, str(df[col].dtype), n_nulos, f"{pct:.2f}%"])
    agregar_tabla(
        doc, ["Variable", "Tipo", "Nulos", "% nulos"], filas_nulos, 4,
        "Tipos de datos y valores nulos por columna (DataFrame transformado).",
        "elaboración propia a partir de AE2/resultados/churn_transformado.csv.",
    )

    h2(doc, "7.3 Distribución de las variables clave")
    parrafo(
        doc,
        f"La variable objetivo churn presenta {churn_no} clientes que permanecen "
        f"(No) y {churn_si} que cancelan (Yes), es decir una tasa de cancelación del "
        f"{pct_churn:.2f}%. Las Tablas 5 a 7 presentan las frecuencias de las "
        f"variables contractuales y de servicio, y la Tabla 8 las estadísticas "
        f"descriptivas de las variables de facturación.",
    )
    agregar_tabla(
        doc, ["contract", "Frecuencia"],
        [[k, int(v)] for k, v in contract_vc.items()], 5,
        "Frecuencia por tipo de contrato.",
        "elaboración propia a partir de AE2/resultados/churn_transformado.csv.",
    )
    agregar_tabla(
        doc, ["internet_service", "Frecuencia"],
        [[k, int(v)] for k, v in internet_vc.items()], 6,
        "Frecuencia por tipo de servicio de internet.",
        "elaboración propia a partir de AE2/resultados/churn_transformado.csv.",
    )
    agregar_tabla(
        doc, ["payment_method", "Frecuencia"],
        [[k, int(v)] for k, v in pago_vc.items()], 7,
        "Frecuencia por método de pago.",
        "elaboración propia a partir de AE2/resultados/churn_transformado.csv.",
    )
    est_filas = []
    for stat in ["count", "mean", "std", "min", "25%", "50%", "75%", "max"]:
        fila = [stat]
        for col in ["tenure", "monthly_charges", "total_charges"]:
            fila.append(f"{desc.loc[stat, col]:.2f}")
        est_filas.append(fila)
    agregar_tabla(
        doc, ["Estadístico", "tenure", "monthly_charges", "total_charges"],
        est_filas, 8, "Estadísticas descriptivas de las variables de facturación.",
        "elaboración propia a partir de AE2/resultados/churn_transformado.csv.",
    )

    h2(doc, "7.4 Evidencia del resumen automático")
    parrafo(
        doc,
        "El resumen se genera automáticamente con la clase DatasetSummarizer y se "
        "guarda en la ruta relativa AE2/resultados/resumen_dataframe.txt, en "
        "codificación UTF-8, regenerándose en cada ejecución del notebook. A "
        "continuación se transcribe su contenido íntegro.",
    )
    parrafo_mono(doc, resumen_txt)

    h2(doc, "7.5 Relación con la necesidad")
    parrafo(
        doc,
        "Las variables churn, contract, payment_method e internet_service, junto con "
        "tenure, monthly_charges y total_charges, respaldan directamente la pregunta "
        "de análisis, pues permiten contrastar el comportamiento contractual, de "
        "servicio y de facturación entre los clientes que cancelan y los que "
        "permanecen. Esta correspondencia confirma la coherencia entre la necesidad "
        "planteada y los datos disponibles.",
    )

    # =========================== 8. CONCLUSIONES ===========================
    h1(doc, "8. Conclusiones")
    for texto in [
        f"El dataset Telco Customer Churn es viable para la necesidad planteada: "
        f"sus {n_c} variables cubren las dimensiones contractual, de servicio y de "
        f"facturación requeridas por la pregunta de análisis.",
        f"La calidad de los datos es adecuada tras la preparación: el único problema "
        f"relevante fueron los {nulos_tc} valores nulos de total_charges, "
        f"identificados y tratados; no se hallaron duplicados.",
        "Existe coherencia entre la pregunta de análisis y las variables "
        "disponibles, en particular la variable objetivo churn y las variables "
        "contractuales, de servicio y de facturación.",
        "La etapa deja preparado el terreno para el análisis exploratorio y el "
        "modelado de las etapas siguientes, con un DataFrame limpio, tipado y "
        "documentado, y un resumen reproducible.",
    ]:
        item(doc, texto)

    # =========================== 9. RECOMENDACIONES ===========================
    h1(doc, "9. Recomendaciones")
    for texto in [
        "Realizar en la etapa 2 un análisis bivariado de churn frente a contract, "
        "tenure y monthly_charges para cuantificar las diferencias entre grupos.",
        "Incorporar visualizaciones (histogramas y gráficos de barras) que faciliten "
        "la comparación de los perfiles de cancelación.",
        "Evaluar la codificación de variables categóricas y el escalado de variables "
        "numéricas de cara al modelado predictivo de la etapa 3.",
        "Documentar y fijar la revisión del dataset para garantizar la "
        "reproducibilidad exacta en futuras ejecuciones.",
    ]:
        item(doc, texto)

    # =========================== 10. REFERENCIAS ===========================
    h1(doc, "10. Referencias")
    referencias = [
        "Han, J., Kamber, M., & Pei, J. (2011). Data mining: Concepts and "
        "techniques (3rd ed.). Morgan Kaufmann.",
        "Lhoest, Q., et al. (2021). Datasets: A community library for natural "
        "language processing. En Proceedings of the 2021 Conference on Empirical "
        "Methods in Natural Language Processing: System Demonstrations (pp. "
        "175-184). Association for Computational Linguistics.",
        "McKinney, W. (2010). Data structures for statistical computing in Python. "
        "En Proceedings of the 9th Python in Science Conference (pp. 56-61).",
        "Provost, F., & Fawcett, T. (2013). Data science for business. O'Reilly "
        "Media.",
        "Verbeke, W., Dejaeger, K., Martens, D., Hur, J., & Baesens, B. (2012). New "
        "insights into churn prediction in the telecommunication sector. European "
        "Journal of Operational Research, 218(1), 211-229.",
        f"scikit-learn. (s. f.). Churn prediction dataset [Conjunto de datos]. "
        f"Hugging Face. {DATASET_URL}",
        "pandas development team. (s. f.). pandas documentation. "
        "https://pandas.pydata.org/docs/",
        "Hugging Face. (s. f.). Datasets documentation. "
        "https://huggingface.co/docs/datasets",
        f"Pedraza Hernández, J. I. (2026). Repositorio del proyecto "
        f"PREICA2602B020038 - Analítica de Datos [Repositorio]. GitHub. {REPO}",
    ]
    for ref in referencias:
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.left_indent = Pt(36)
        p.paragraph_format.first_line_indent = Pt(-36)
        p.add_run(ref)

    doc.save(str(SALIDA))
    print(f"Documento generado: {SALIDA}")
    print(f"  crudo: {n_crudo_f}x{n_crudo_c} | transformado: {n_f}x{n_c}")
    print(f"  nulos total_charges (crudo): {nulos_tc} | churn Yes: {churn_si} "
          f"({pct_churn:.2f}%)")


if __name__ == "__main__":
    construir()
